"""
Salida: un Excel por proyecto con identidad GYF y, de cada pestaña, un CSV plano.
Las pestañas se construyen una vez (construir_pestanas) y de ahí salen ambos.

Orden de pestañas: Lectura (lo primero que se mira) → Keywords → Familias →
Preguntas → Competidores → Semillas → Descartadas → Avisos → Briefing.
"""

import csv
import io
import os
from collections import Counter, defaultdict
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelos import Aviso, Briefing, PaginaCompetidor, Pregunta, Semilla
from .tabla import Tabla
from .volumen import FUENTE_GSC, SIN_DATO

FUCSIA = "E0007A"
NEGRO = "111111"
GRIS = "777777"
FILA_ALTERNA = "FBEFF5"
SECCION_BG = "3D001A"
SECCION_FG = "FF4DAA"
LOGO = os.path.join(os.path.dirname(__file__), "..", "assets", "logo_gyf.png")

# nombre → (columnas, filas, anchos, nota bajo el título)
Pestanas = Dict[str, Tuple[List[str], List[list], List[int], str]]

NOTA_FAMILIAS = ("VISTA DE LECTURA · BORRADOR. Agrupa por servicio + zona + intención para leer menos filas. "
                 "No es clustering ni decide páginas: eso es de Nuria.")
NOTA_HUECO = ("«hueco = sí»: algún competidor tiene página de ese servicio y zona (o entrada de blog parecida) "
              "y la web propia no. «inferencia débil»: la página del competidor solo se reconoce por el title o la URL.")


def _vol(kw):
    return kw.volumen if kw.volumen is not None else (kw.volumen_texto or None)


def _lectura(briefing, tabla, semillas, paginas, avisos, preguntas, extra) -> List[list]:
    kws = list(tabla.kws.values())
    n = len(kws) or 1
    filas: List[list] = []

    def bloque(nombre, pares):
        for concepto, valor, detalle in pares:
            filas.append([nombre, concepto, valor, detalle])

    bloque("Recogida", [
        ("Fecha", extra["fecha"], ""),
        ("Duración", f"{extra['duracion_seg'] // 60} min {extra['duracion_seg'] % 60} s", ""),
        ("Keywords únicas", len(kws), "después de quitar el ruido"),
        ("Descartadas", len(tabla.descartadas), "ver pestaña Descartadas"),
        ("Semillas", len(semillas), f"{sum(s.prioritaria for s in semillas)} prioritarias"),
        ("Consultas al autocomplete", extra["consultas"], extra["parado"] or "completo dentro del tiempo"),
        ("Dominios espejados", len({p.dominio for p in paginas}), f"{len(paginas)} páginas"),
        ("Preguntas (PAA y relacionadas)", len(preguntas), extra["preguntas_estado"]),
        ("Coste DataForSEO", extra["coste_dfs"], "por consulta, no por suscripción"),
    ])
    for nombre, campo in (("Por intención", "intencion"), ("Por servicio", "servicio")):
        c = Counter(getattr(k, campo) or "(sin servicio)" for k in kws)
        bloque(nombre, [(k, v, f"{v * 100 // n} %") for k, v in c.most_common()])
    zonas = Counter()
    fuera = Counter()
    for k in kws:
        if k.geo.endswith("(fuera de briefing)"):
            zonas["otras zonas (fuera de briefing)"] += 1
            fuera[k.geo.replace(" (fuera de briefing)", "")] += 1
        else:
            zonas[k.geo or "(sin zona)"] += 1
    bloque("Por zona", [(k, v, f"{v * 100 // n} %") for k, v in zonas.most_common()])
    bloque("Zonas fuera de briefing más nombradas", [(k, v, "¿añadir al briefing?") for k, v in fuera.most_common(15)])
    modulos = Counter(m for k in kws for m in k.modulos)
    bloque("Por módulo de origen", [(k, v, "una keyword puede venir de varios") for k, v in modulos.most_common()])
    bloque("Volumen por fuente", [(k, v, "") for k, v in extra["volumen_por_fuente"].items()])

    gsc = sorted([k for k in kws if k.fuente_volumen == FUENTE_GSC and k.intencion != "marca"],
                 key=lambda k: -(k.volumen or 0))[:20]
    bloque("Top Search Console (sin marca)", [
        (k.texto, k.volumen, f"posición {k.posicion_media} · {int(k.clics or 0)} clics · {k.intencion}") for k in gsc])

    fam = Counter(k.familia for k in kws)
    bloque("Familias más grandes (borrador)", [(f, v, "") for f, v in fam.most_common(15)])

    huecos = defaultdict(lambda: [0, set(), set()])
    for k in kws:
        if k.hueco == "sí":
            h = huecos[k.familia]
            h[0] += 1
            h[1].update(k.competidores_con_pagina.split(", "))
            h[2].add(k.fiabilidad_hueco)
    top_huecos = sorted(huecos.items(), key=lambda x: (-len(x[1][1]), -x[1][0]))[:15]
    bloque("Huecos frente a la competencia", [
        (f, v[0], f"{', '.join(sorted(v[1]))} · {'fuerte' if 'fuerte' in v[2] else 'inferencia débil'}")
        for f, v in top_huecos])

    fiab = defaultdict(Counter)
    for p in paginas:
        if p.fiabilidad != "no aplica":
            fiab[p.dominio][p.fiabilidad] += 1
    bloque("Fiabilidad del espejo por dominio", [
        (d, sum(c.values()), " · ".join(f"{k}: {v}" for k, v in c.most_common())) for d, c in fiab.items()])

    bloque("Descartadas por motivo", [(k, v, "") for k, v in
                                      Counter(d.motivo for d in tabla.descartadas.values()).most_common()])
    bloque("Avisos por tipo", [(k, v, "") for k, v in Counter(a.tipo for a in avisos).most_common()])
    return filas


def _familias(tabla: Tabla) -> List[list]:
    grupos = defaultdict(list)
    for k in tabla.kws.values():
        grupos[k.familia].append(k)
    filas = []
    for fam, ks in grupos.items():
        k0 = ks[0]
        geo = "otras zonas" if k0.geo.endswith("(fuera de briefing)") else (k0.geo or "sin zona")
        imp = sum(k.volumen or 0 for k in ks if k.fuente_volumen == FUENTE_GSC)
        busq = sum(k.volumen or 0 for k in ks if k.fuente_volumen not in (FUENTE_GSC, SIN_DATO)
                   and k.volumen is not None)
        con_hueco = [k for k in ks if k.hueco == "sí"]
        comps = sorted({c for k in con_hueco for c in k.competidores_con_pagina.split(", ") if c})
        propia = next((k.pagina_propia for k in ks if k.pagina_propia), "")
        ejemplos = " | ".join(k.texto for k in sorted(ks, key=lambda k: -k.apariciones)[:5])
        filas.append([fam, k0.servicio or "sin servicio", geo, k0.intencion, len(ks),
                      imp or None, busq or None, len(con_hueco) or None, ", ".join(comps), propia, ejemplos])
    orden = {"transaccional": 0, "local": 1, "marca": 2, "comparativa": 3, "informacional": 4}
    return sorted(filas, key=lambda f: (orden.get(f[3], 9), -f[4]))


def construir_pestanas(briefing: Briefing, tabla: Tabla, semillas: List[Semilla],
                       paginas: List[PaginaCompetidor], avisos: List[Aviso],
                       preguntas: List[Pregunta], extra: dict) -> Pestanas:
    kws = [[kw.texto, kw.semilla_origen, ", ".join(kw.modulos), kw.intencion, kw.geo, kw.servicio,
            kw.familia, _vol(kw), kw.fuente_volumen, kw.marca_estimacion, kw.apariciones,
            kw.clics, kw.posicion_media, kw.cpc, kw.hueco, kw.competidores_con_pagina,
            kw.pagina_propia, kw.fiabilidad_hueco]
           for kw in tabla.ordenadas()]
    comp = [[p.dominio, p.url, p.title, p.h1, p.meta_description, p.h2, p.patron, p.keyword_inferida,
             p.servicio, p.geo, p.fiabilidad] for p in paginas]
    sem = [[s.texto, s.servicio, s.intencion, s.geo, s.tipo, "sí" if s.prioritaria else "",
            s.sugerencias, s.consultas, s.estado] for s in semillas]
    preg = [[p.pregunta, p.tipo, p.semilla, p.servicio, p.geo] for p in preguntas]
    desc = [[d.keyword, d.motivo, d.termino, d.modulo, d.semilla, d.veces]
            for d in sorted(tabla.descartadas.values(), key=lambda d: (d.motivo, d.termino, d.keyword))]
    avi = [[a.tipo, a.elemento, a.detalle] for a in avisos]

    b = briefing
    brief = [
        ["Dominio", b.dominio], ["Marca", b.marca], ["Sector / categoría", b.sector],
        ["Servicios núcleo", "\n".join(("* " if s.prioritario else "") + s.nombre for s in b.servicios)],
        ["Lo que NO ofrece (lista negra del cliente)", "\n".join(b.no_ofrece)],
        ["Lista negra por defecto (ruido)", "de fábrica" if b.ruido is None else ", ".join(b.ruido) or "desactivada"],
        ["Objetivo", b.objetivo], ["Sede", b.sede], ["Zona principal", b.zona_principal],
        ["Zonas secundarias", "\n".join(b.zonas_secundarias)], ["Alcance", b.alcance],
        ["Web", b.web_estado + (f" · {b.url_web}" if b.url_web else "")],
        ["Sitemap propio", b.sitemap_url or ""],
        ["Export de Search Console", b.gsc_nombre or ("subido" if b.gsc_csv else "no")],
        ["Export de Keyword Planner", b.planner_nombre or ("subido" if b.planner_csv else "no")],
        ["Competidores de referencia", "\n".join(b.competidores)],
        ["Jerga del público", "\n".join(b.jerga)], ["Idioma / mercado", b.mercado],
        ["Segundo nivel de sopa de letras", "sí" if b.segundo_nivel else "no"],
        ["Preguntas (PAA) con DataForSEO", "sí" if b.preguntas else "no"],
        ["Volumen con DataForSEO", "sí" if b.volumen_dataforseo else "no"],
        ["Modificadores", "\n".join(f"{k}: {', '.join(v)}" for k, v in (b.modificadores or {}).items())
         or "por defecto"],
        ["Fecha de la recogida", extra["fecha"]],
    ]

    return {
        "Lectura": (["bloque", "concepto", "valor", "detalle"],
                    _lectura(briefing, tabla, semillas, paginas, avisos, preguntas, extra),
                    [30, 46, 14, 70], "Los recuentos del informe, ya hechos. Lo primero que se mira."),
        "Keywords": (["keyword", "semilla origen", "módulo", "intención", "geo detectada",
                      "servicio detectado", "familia (borrador)", "volumen", "fuente del volumen",
                      "marca de estimación", "nº de apariciones", "clics (GSC)", "posición media (GSC)",
                      "CPC (DataForSEO)", "hueco", "competidores con página", "página propia",
                      "fiabilidad del hueco"],
                     kws, [38, 26, 20, 14, 20, 20, 34, 12, 26, 14, 11, 10, 11, 10, 8, 30, 36, 15], NOTA_HUECO),
        "Familias": (["familia", "servicio", "zona", "intención", "nº keywords", "impresiones (GSC)",
                      "búsquedas/mes (Planner o DataForSEO)", "keywords con hueco", "competidores con página",
                      "página propia", "ejemplos"],
                     _familias(tabla), [40, 20, 18, 14, 11, 13, 15, 12, 30, 36, 80], NOTA_FAMILIAS),
        "Preguntas": (["pregunta", "tipo", "semilla", "servicio", "geo"], preg, [60, 22, 34, 22, 18],
                      "People Also Ask y búsquedas relacionadas: materia del blog. Intención informacional por defecto."),
        "Competidores": (["dominio", "URL", "title", "H1", "meta description", "H2", "patrón de URL",
                          "keyword inferida", "servicio", "zona", "fiabilidad"],
                         comp, [22, 48, 40, 34, 50, 50, 15, 32, 20, 18, 16], ""),
        "Semillas": (["semilla", "servicio", "intención", "geo", "tipo", "prioritaria",
                      "nº de sugerencias que produjo", "consultas hechas", "estado"],
                     sem, [38, 24, 15, 16, 12, 11, 14, 12, 14], ""),
        "Descartadas": (["keyword", "motivo", "término", "módulo", "semilla", "veces"], desc,
                        [44, 30, 16, 16, 30, 8],
                        "Lo que quitó el filtro de ruido, con su motivo, por si hay que rescatar algo."),
        "Avisos": (["tipo", "elemento afectado", "detalle"], avi, [22, 50, 90], ""),
        "Briefing": (["campo", "valor"], brief, [36, 80], "Copia del formulario, para trazabilidad y para relanzar."),
    }


def _escribir_hoja(ws, titulo: str, subtitulo: str, nota: str, columnas, filas, anchos, secciones=False):
    n = len(columnas)
    ultima = get_column_letter(n)
    ws.sheet_view.showGridLines = False

    # Cabecera de marca: logo + título en fucsia
    ws.row_dimensions[1].height = 52
    if os.path.exists(LOGO):
        img = XLImage(LOGO)
        img.width, img.height = img.width * 0.62, img.height * 0.62
        ws.add_image(img, "A1")
    ws.merge_cells(f"A2:{ultima}2")
    ws["A2"] = titulo
    ws["A2"].font = Font(name="Arial", size=13, bold=True, color=FUCSIA)
    ws.merge_cells(f"A3:{ultima}3")
    ws["A3"] = subtitulo
    ws["A3"].font = Font(name="Arial", size=9, color=GRIS)
    ws.merge_cells(f"A4:{ultima}4")
    ws["A4"] = nota
    ws["A4"].font = Font(name="Arial", size=9, italic=True, color=FUCSIA)

    fila_cab = 6
    for i, col in enumerate(columnas, 1):
        c = ws.cell(row=fila_cab, column=i, value=col)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=FUCSIA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila_cab].height = 30

    alterna = PatternFill("solid", fgColor=FILA_ALTERNA)
    seccion = PatternFill("solid", fgColor=SECCION_BG)
    r = fila_cab + 1
    actual = None
    for fila in filas:
        if secciones and fila[0] != actual:
            actual = fila[0]
            for i in range(1, n + 1):
                ws.cell(row=r, column=i).fill = seccion
            c = ws.cell(row=r, column=1, value=f"  {actual}")
            c.font = Font(name="Arial", size=11, bold=True, color=SECCION_FG)
            r += 1
        for i, v in enumerate(fila, 1):
            c = ws.cell(row=r, column=i, value=v if v != "" else None)
            c.font = Font(name="Arial", size=10, color=NEGRO)
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if (r - fila_cab) % 2 == 0:
                c.fill = alterna
        r += 1

    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=1)
    if filas and not secciones:
        ws.auto_filter.ref = f"A{fila_cab}:{ultima}{fila_cab + len(filas)}"


def generar_excel(ruta: str, briefing: Briefing, pestanas: Pestanas, resumen: str):
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, (cols, filas, anchos, nota) in pestanas.items():
        ws = wb.create_sheet(nombre)
        _escribir_hoja(ws, f"LA LLAVE · {briefing.marca.upper()} · {nombre.upper()}",
                       f"{resumen}  ·  El Gordo y el Flaco Marketing Online", nota, cols, filas, anchos,
                       secciones=(nombre == "Lectura"))

    # Regla que no se toca: ningún volumen sin fuente
    ws = wb["Keywords"]
    for fila in ws.iter_rows(min_row=7, min_col=8, max_col=9):
        vol, fuente = fila
        if vol.value is not None:
            assert fuente.value and fuente.value != SIN_DATO, "volumen sin fuente"
            if isinstance(vol.value, (int, float)):
                vol.number_format = "#,##0"
    wb.save(ruta)


def generar_csv(pestanas: Pestanas, nombre: str) -> str:
    cols, filas, _, _ = pestanas[nombre]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for f in filas:
        w.writerow(["" if v is None else v for v in f])
    return buf.getvalue()


def a_json(pestanas: Pestanas) -> dict:
    return {n: {"columnas": c, "filas": f, "anchos": a, "nota": nota} for n, (c, f, a, nota) in pestanas.items()}


def de_json(datos: dict) -> Pestanas:
    return {n: (d["columnas"], d["filas"], d["anchos"], d["nota"]) for n, d in datos.items()}

